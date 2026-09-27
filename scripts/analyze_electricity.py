#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""月度电费分析脚本 - 读取南网供电电费清单Excel，输出结构化分析数据"""

import argparse
import json
import os
import sys
import math
from pathlib import Path

try:
    import pandas as pd
    import openpyxl
except ImportError:
    print("ERROR: pandas and openpyxl are required. Install with: pip install pandas openpyxl")
    sys.exit(1)


def safe_float(val):
    if val is None or (isinstance(val, float) and math.isnan(val)):
        return 0.0
    return float(val)


def parse_percent(val):
    if pd.isna(val):
        return 0.0
    if isinstance(val, str):
        return float(val.rstrip('%'))
    return float(val)


def analyze_workbook(filepath, month_label):
    """Read and analyze the Excel workbook"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    
    # Check required sheets
    required = ['主要分析结果', '异常用户明细', '电表用途统计', '基准值统计']
    for s in required:
        if s not in wb.sheetnames:
            print(f"WARNING: Sheet '{s}' not found. Available: {wb.sheetnames}")
    
    # 主要分析结果 has a title row in row 1, actual headers in row 2 (header=1)
    df_main = pd.read_excel(filepath, sheet_name='主要分析结果', header=1)
    df_abnormal = pd.read_excel(filepath, sheet_name='异常用户明细')
    df_usage = pd.read_excel(filepath, sheet_name='电表用途统计')
    df_baseline = pd.read_excel(filepath, sheet_name='基准值统计')
    
    result = {
        'month_label': month_label,
        'generated_at': pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S'),
        'source_file': filepath,
    }
    
    # === 1. Overall Statistics ===
    total_current = safe_float(df_main['当前月应收电费'].sum())
    total_last = safe_float(df_main['上月应收电费'].sum())
    total_diff = total_current - total_last
    total_growth_rate = (total_diff / total_last * 100) if total_last > 0 else 0
    
    status_counts = df_main['用户状态'].fillna('未标注').value_counts().to_dict()
    
    result['overall'] = {
        'total_current': round(total_current, 1),
        'total_last': round(total_last, 1),
        'total_diff': round(total_diff, 1),
        'growth_rate': round(total_growth_rate, 1),
        'total_users': len(df_main),
        'status': {k: int(v) for k, v in status_counts.items()},
    }
    
    # === 2. Usage Category Analysis ===
    usage_categories = ['防溺水监控', '接入网机房', '综合机房', '治安视频监控']
    usage_summary = []
    for purpose in usage_categories:
        sub = df_main[df_main['电表用途'] == purpose]
        curr = safe_float(sub['当前月应收电费'].sum())
        last = safe_float(sub['上月应收电费'].sum())
        diff = curr - last
        pct = (diff / last * 100) if last > 0 else 0
        share = (curr / total_current * 100) if total_current > 0 else 0
        usage_summary.append({
            'purpose': purpose,
            'user_count': len(sub),
            'current_fee': round(curr, 1),
            'last_fee': round(last, 1),
            'diff': round(diff, 1),
            'growth_rate': round(pct, 1),
            'share': round(share, 1),
        })
    
    # Unlabeled
    sub_unlabeled = df_main[df_main['电表用途'].isna()]
    curr_u = safe_float(sub_unlabeled['当前月应收电费'].sum())
    last_u = safe_float(sub_unlabeled['上月应收电费'].sum())
    diff_u = curr_u - last_u
    pct_u = (diff_u / last_u * 100) if last_u > 0 else 0
    share_u = (curr_u / total_current * 100) if total_current > 0 else 0
    usage_summary.append({
        'purpose': '用途待核实',
        'user_count': len(sub_unlabeled),
        'current_fee': round(curr_u, 1),
        'last_fee': round(last_u, 1),
        'diff': round(diff_u, 1),
        'growth_rate': round(pct_u, 1),
        'share': round(share_u, 1),
    })
    result['usage_summary'] = usage_summary
    
    # === 3. Meter Usage Statistics (from台账) ===
    meter_usage = []
    for _, row in df_usage.iterrows():
        meter_usage.append({
            'purpose': str(row['电表用途']),
            'count': int(row['数量']) if pd.notna(row['数量']) else 0,
            'share': str(row['占比']) if pd.notna(row['占比']) else '',
        })
    result['meter_usage'] = meter_usage
    
    # === 4. Fee Distribution ===
    bins = [0, 10, 50, 100, 500, 1000, 5000, 10000, 50000, 999999999]
    labels = ['0-10', '10-50', '50-100', '100-500', '500-1000', '1000-5000', '5000-10000', '10000-50000', '50000+']
    df_main['_fee_bin'] = pd.cut(df_main['当前月应收电费'], bins=bins, labels=labels, right=False)
    fee_dist = df_main['_fee_bin'].value_counts().sort_index()
    result['fee_distribution'] = [
        {'range': label, 'count': int(fee_dist.get(label, 0))} for label in labels
    ]
    
    # === 5. Abnormal Users ===
    abnormal_list = []
    for _, row in df_abnormal.iterrows():
        abnormal_list.append({
            'user_id': str(row['用户编号']),
            'address': str(row['用电地址']),
            'current_fee': safe_float(row['当前月应收电费']),
            'last_fee': safe_float(row['上月应收电费']),
            'diff': safe_float(row['电费差值']),
            'growth_rate': str(row['电费增长率']),
            'baseline_growth': str(row['相对于基准增长率']),
            'baseline_growth_num': parse_percent(row['相对于基准增长率']),
            'purpose': str(row['电表用途']) if pd.notna(row['电表用途']) else '',
        })
    
    # Sort by diff descending
    abnormal_list.sort(key=lambda x: x['diff'], reverse=True)
    
    # Classify by severity
    severe = [u for u in abnormal_list if u['baseline_growth_num'] >= 50]
    moderate = [u for u in abnormal_list if 20 <= u['baseline_growth_num'] < 50]
    
    result['abnormal'] = {
        'total_count': len(abnormal_list),
        'total_current_fee': round(sum(u['current_fee'] for u in abnormal_list), 1),
        'total_last_fee': round(sum(u['last_fee'] for u in abnormal_list), 1),
        'total_diff': round(sum(u['diff'] for u in abnormal_list), 1),
        'severe_count': len(severe),
        'moderate_count': len(moderate),
        'top20': abnormal_list[:20],
        'severe_top10': sorted(severe, key=lambda x: x['diff'], reverse=True)[:10],
    }
    
    # === 6. Abnormal by Purpose ===
    abnormal_by_purpose = []
    for purpose in ['防溺水监控', '接入网机房', '综合机房', '治安视频监控', '']:
        if purpose:
            sub = [u for u in abnormal_list if u['purpose'] == purpose]
            label = purpose
        else:
            sub = [u for u in abnormal_list if u['purpose'] == '']
            label = '用途待核实'
        if sub:
            abnormal_by_purpose.append({
                'purpose': label,
                'count': len(sub),
                'current_fee': round(sum(u['current_fee'] for u in sub), 1),
                'last_fee': round(sum(u['last_fee'] for u in sub), 1),
                'diff': round(sum(u['diff'] for u in sub), 1),
            })
    result['abnormal_by_purpose'] = abnormal_by_purpose
    
    # === 7. Top20 High Consumption Users ===
    df_top = df_main.sort_values('当前月应收电费', ascending=False).head(20)
    top_consumers = []
    for _, row in df_top.iterrows():
        top_consumers.append({
            'address': str(row['用电地址']),
            'current_fee': safe_float(row['当前月应收电费']),
            'last_fee': safe_float(row['上月应收电费']),
            'purpose': str(row['电表用途']) if pd.notna(row['电表用途']) else '',
            'diff': round(safe_float(row['当前月应收电费']) - safe_float(row['上月应收电费']), 1),
        })
    result['top_consumers'] = top_consumers
    
    # === 8. New and Lost Users ===
    new_users = []
    for _, row in df_main[df_main['用户状态'] == '新增'].iterrows():
        new_users.append({
            'address': str(row['用电地址']),
            'current_fee': safe_float(row['当前月应收电费']),
            'purpose': str(row['电表用途']) if pd.notna(row['电表用途']) else '',
        })
    
    lost_users = []
    for _, row in df_main[df_main['用户状态'] == '流失'].iterrows():
        lost_users.append({
            'address': str(row['用电地址']),
            'last_fee': safe_float(row['上月应收电费']),
            'purpose': str(row['电表用途']) if pd.notna(row['电表用途']) else '',
        })
    
    result['new_users'] = new_users
    result['lost_users'] = lost_users
    
    return result


def load_history(history_dir):
    """Load all historical analysis JSON files"""
    history = []
    if not os.path.exists(history_dir):
        return history
    for f in sorted(os.listdir(history_dir)):
        if f.startswith('analysis_') and f.endswith('.json'):
            try:
                with open(os.path.join(history_dir, f), 'r', encoding='utf-8') as fh:
                    data = json.load(fh)
                    history.append(data)
            except Exception as e:
                print(f"WARNING: Failed to load {f}: {e}")
    return history


def compare_with_history(current, history):
    """Compare current analysis with previous month's data"""
    if not history:
        return None
    # Find the most recent previous analysis
    prev = history[-1]
    comparison = {
        'prev_month': prev.get('month_label', ''),
        'prev_total': prev.get('overall', {}).get('total_current', 0),
        'prev_abnormal_count': prev.get('abnormal', {}).get('total_count', 0),
        'prev_usage': {u['purpose']: u for u in prev.get('usage_summary', [])},
    }
    return comparison


def main():
    parser = argparse.ArgumentParser(description='月度电费分析')
    parser.add_argument('--input', required=True, help='Excel文件路径')
    parser.add_argument('--output', required=True, help='输出目录')
    parser.add_argument('--month', required=True, help='分析月份标签，如"2026年8月"')
    parser.add_argument('--history-dir', default='', help='历史分析数据目录')
    args = parser.parse_args()
    
    filepath = args.input
    if not os.path.exists(filepath):
        print(f"ERROR: File not found: {filepath}")
        sys.exit(1)
    
    os.makedirs(args.output, exist_ok=True)
    
    print(f"分析月份: {args.month}")
    print(f"数据源: {filepath}")
    print(f"输出目录: {args.output}")
    print()
    
    result = analyze_workbook(filepath, args.month)
    
    # Load history for comparison
    history_dir = args.history_dir or os.path.join(os.path.dirname(os.path.dirname(filepath)), '历史分析数据')
    history = load_history(history_dir)
    if history:
        comparison = compare_with_history(result, history)
        if comparison:
            result['comparison'] = comparison
            print(f"已加载历史数据 {len(history)} 期，最近一期: {comparison['prev_month']}")
    else:
        print("未找到历史分析数据（首次分析）")
    
    # Save result JSON
    output_file = os.path.join(args.output, 'analysis_result.json')
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print(f"\n分析结果已保存: {output_file}")
    
    # Print summary
    ov = result['overall']
    print(f"\n{'='*60}")
    print(f"  {args.month} 电费分析摘要")
    print(f"{'='*60}")
    print(f"  总电费: {ov['total_current']:.1f}元 (上月: {ov['total_last']:.1f}元)")
    print(f"  环比: {'+' if ov['total_diff']>=0 else ''}{ov['total_diff']:.1f}元 ({ov['growth_rate']:.1f}%)")
    print(f"  用户数: {ov['total_users']} (正常{ov['status'].get('正常',0)} 新增{ov['status'].get('新增',0)} 流失{ov['status'].get('流失',0)})")
    print(f"  异常用户: {result['abnormal']['total_count']}户 (严重{result['abnormal']['severe_count']} 中度{result['abnormal']['moderate_count']})")
    print(f"  异常增量: {result['abnormal']['total_diff']:.1f}元")
    print(f"\n  按用途分类:")
    for u in result['usage_summary']:
        print(f"    {u['purpose']}: {u['user_count']}户, {u['current_fee']:.1f}元 ({u['growth_rate']:+.1f}%)")
    print(f"{'='*60}")


if __name__ == '__main__':
    main()
